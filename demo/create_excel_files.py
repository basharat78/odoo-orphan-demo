"""
Generate orphans_data.xlsx and donors_data.xlsx for the demo.
Run once to create the sample Excel files.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── colour palette ───────────────────────────────────────────────────────────
GREEN  = "1F7A4A"   # header bg  (orphans sheet)
BLUE   = "1A3C6E"   # header bg  (donors sheet)
LGREY  = "F2F2F2"   # alt row
WHITE  = "FFFFFF"
GOLD   = "F5C518"   # accent

def header_font():   return Font(name="Calibri", bold=True, color=WHITE, size=11)
def data_font():     return Font(name="Calibri", size=11)
def center():        return Alignment(horizontal="center", vertical="center")
def left():          return Alignment(horizontal="left",   vertical="center")
def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def style_sheet(ws, col_widths, header_color):
    """Apply header styling and column widths."""
    fill = PatternFill("solid", fgColor=header_color)
    for cell in ws[1]:
        cell.font      = header_font()
        cell.fill      = fill
        cell.alignment = center()
        cell.border    = thin_border()
    ws.row_dimensions[1].height = 28

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        bg = LGREY if row_idx % 2 == 0 else WHITE
        fill_row = PatternFill("solid", fgColor=bg)
        for cell in row:
            cell.font      = data_font()
            cell.fill      = fill_row
            cell.border    = thin_border()
            cell.alignment = center() if cell.column > 1 else left()
    ws.freeze_panes = "A2"


# ────────────────────────────────────────────────────────────────────────────
#  ORPHANS workbook
# ────────────────────────────────────────────────────────────────────────────
orphans = [
    # name,               age, gender,   location,    notes
    ("Ahmed Ali",          7,  "male",   "Karachi",   "Loves football"),
    ("Fatima Hassan",      9,  "female", "Lahore",    "Top student in class"),
    ("Yusuf Khan",         5,  "male",   "Islamabad", "Very energetic child"),
    ("Aisha Malik",       11,  "female", "Peshawar",  "Dreams of becoming a doctor"),
    ("Omar Siddiqui",      8,  "male",   "Multan",    "Enjoys drawing"),
    ("Zainab Qureshi",     6,  "female", "Quetta",    "Loves to sing"),
    ("Hassan Raza",       10,  "male",   "Faisalabad","Excellent at mathematics"),
    ("Maryam Chaudhry",   12,  "female", "Hyderabad", "Wants to be a teacher"),
    ("Abdullah Sheikh",    4,  "male",   "Rawalpindi","Youngest in the program"),
    ("Noor Bibi",          8,  "female", "Sialkot",   "Kind and helpful"),
]

wb_orphans = openpyxl.Workbook()
ws = wb_orphans.active
ws.title = "Orphans"

ws.append(["Full Name", "Age", "Gender", "City / Location", "Notes"])
for row in orphans:
    ws.append(list(row))

style_sheet(ws, col_widths=[22, 8, 12, 18, 32], header_color=GREEN)

wb_orphans.save("/Users/mac/odoo/orphans_data.xlsx")
print("✓  orphans_data.xlsx  created  (10 orphans)")


# ────────────────────────────────────────────────────────────────────────────
#  DONORS workbook
# ────────────────────────────────────────────────────────────────────────────
donors = [
    # name,              email,                              phone
    ("Waleed Malik",     "mbasharat596@gmail.com",           "+92-300-1234567"),
    ("Sara Ahmed",       "popen9488@gmail.com",              "+92-321-2345678"),
    ("Bilal Hussain",    "waleed.malik5136434@gmail.com",    "+92-333-3456789"),
    ("Amina Farooq",     "mbasharat596@gmail.com",           "+92-345-4567890"),
    ("Khalid Mehmood",   "popen9488@gmail.com",              "+92-312-5678901"),
]

wb_donors = openpyxl.Workbook()
ws2 = wb_donors.active
ws2.title = "Donors"

ws2.append(["Full Name", "Email", "Phone"])
for row in donors:
    ws2.append(list(row))

style_sheet(ws2, col_widths=[22, 30, 20], header_color=BLUE)

wb_donors.save("/Users/mac/odoo/donors_data.xlsx")
print("✓  donors_data.xlsx   created  (5 donors)")
